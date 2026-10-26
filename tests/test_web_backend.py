"""Business rules of the web backend (fake neural network, real pipeline + SQLite).

Run from the project root:   attendance_system\\.venv\\Scripts\\python -m pytest tests
"""
import io
import os
import sqlite3
import threading
import unicodedata
from datetime import datetime, timedelta

import numpy as np
import pytest
from openpyxl import load_workbook

from backend import clock
from conftest import make_container, unit


# ---------------------------------------------------------------- classes/courses
def test_create_class_and_course_normalised(client):
    r = client.post('/api/classes', json={'name': '  se1501 ', 'description': 'Fall'})
    assert r.status_code == 201 and r.json()['name'] == 'SE1501'
    assert os.path.isdir(os.path.join(client.container.paths.dataset_dir, 'SE1501'))
    r = client.post('/api/courses', json={'class_name': 'se1501', 'subject': ' cpv301 '})
    assert r.status_code == 201 and r.json()['subject'] == 'CPV301'
    assert client.post('/api/classes', json={'name': 'SE1501'}).status_code == 409
    assert client.post('/api/courses', json={'class_name': 'SE1501', 'subject': 'CPV301'}).status_code == 409


@pytest.mark.parametrize('bad', ['../evil', 'a_b', 'SE 1501', '..', 'CON', 'x/y', '', 'a' * 40])
def test_invalid_class_names_rejected(client, bad):
    r = client.post('/api/classes', json={'name': bad})
    assert r.status_code == 422, bad
    assert not os.path.exists(os.path.join(client.container.paths.data_dir, 'evil'))


def test_delete_course_keeps_history(h, client):
    h.course('SE1', 'CPV')
    _, emb = h.enroll('SE1', 'S1', 'An')
    sid = h.open('SE1', 'CPV').json()['id']
    assert h.frame(sid, h.face_frame(emb))['status'] == 'recognized'
    course_id = client.get('/api/courses').json()[0]['id']
    assert client.delete(f'/api/courses/{course_id}').status_code == 200
    hist = client.get('/api/history').json()
    assert hist['summary']['present'] == 1
    # class deletion preserves attendance history
    cls_id = client.get('/api/classes').json()[0]['id']
    assert client.delete(f'/api/classes/{cls_id}').status_code in (200, 409)
    assert client.get('/api/history').json()['summary']['present'] == 1


# ---------------------------------------------------------------- sessions
def test_two_classes_same_subject_slot_day_are_separate(h, client):
    h.course('A1', 'CPV')
    h.course('B1', 'CPV')
    _, ea = h.enroll('A1', 'S1', 'An')
    _, eb = h.enroll('B1', 'S1', 'Bình')      # same Student ID in another class
    sa = h.open('A1', 'CPV', slot=2).json()
    sb = h.open('B1', 'CPV', slot=2).json()
    assert sa['id'] != sb['id']
    assert h.frame(sa['id'], h.face_frame(ea))['status'] == 'recognized'
    # B1's student is not in A1's gallery → stranger in A1's session
    assert h.frame(sa['id'], h.face_frame(eb))['status'] == 'unknown'
    assert h.frame(sb['id'], h.face_frame(eb))['status'] == 'recognized'
    ra = client.get(f"/api/sessions/{sa['id']}/records").json()['records']
    rb = client.get(f"/api/sessions/{sb['id']}/records").json()['records']
    assert [(r['student_name'], r['status']) for r in ra] == [('An', 'present')]
    assert [(r['student_name'], r['status']) for r in rb] == [('Bình', 'present')]


def test_two_subjects_one_class_same_slot_day_are_separate(h, client):
    h.course('A1', 'CPV')
    h.course('A1', 'MAD')
    _, e = h.enroll('A1', 'S1', 'An')
    s1 = h.open('A1', 'CPV', slot=3).json()
    s2 = h.open('A1', 'MAD', slot=3).json()
    assert s1['id'] != s2['id']
    assert h.frame(s1['id'], h.face_frame(e))['status'] == 'recognized'
    assert h.frame(s2['id'], h.face_frame(e))['status'] == 'recognized'
    with sqlite3.connect(client.container.paths.db_path) as db:
        rows = db.execute('SELECT subject, session_id FROM attendance ORDER BY subject').fetchall()
    assert rows == [('CPV', s1['id']), ('MAD', s2['id'])]


def test_reopen_keeps_records_and_timestamps(h, client):
    h.course('A1', 'CPV')
    _, e = h.enroll('A1', 'S1', 'An')
    s = h.open('A1', 'CPV').json()
    first = h.frame(s['id'], h.face_frame(e))['attendance']['timestamp']
    client.post(f"/api/sessions/{s['id']}/end")
    r = h.open('A1', 'CPV')
    assert r.status_code == 409 and r.json()['code'] == 'session_ended'
    r = h.open('A1', 'CPV', reopen=True)
    assert r.status_code == 201 and r.json()['id'] == s['id'] and r.json()['reopened']
    again = h.frame(s['id'], h.face_frame(e))
    assert again['status'] == 'already_recorded' and again['attendance']['timestamp'] == first


def test_reload_does_not_reset_session(h, client):
    h.course('A1', 'CPV')
    a = h.open('A1', 'CPV', minutes=30).json()
    b = h.open('A1', 'CPV', minutes=5).json()          # same page reloaded / opened twice
    assert a['id'] == b['id'] and a['deadline_at'] == b['deadline_at'] and not b['created']
    assert client.get(f"/api/sessions/{a['id']}").json()['deadline_at'] == a['deadline_at']


def test_recognised_many_times_is_one_row(h, client):
    h.course('A1', 'CPV')
    _, e = h.enroll('A1', 'S1', 'An')
    s = h.open('A1', 'CPV').json()
    statuses = [h.frame(s['id'], h.face_frame(e), frame_id=i)['status'] for i in range(5)]
    assert statuses == ['recognized'] + ['already_recorded'] * 4
    with sqlite3.connect(client.container.paths.db_path) as db:
        assert db.execute('SELECT COUNT(*) FROM attendance').fetchone()[0] == 1


def test_concurrent_writes_insert_once(h, client):
    h.course('A1', 'CPV')
    h.enroll('A1', 'S1', 'An')
    sid = h.open('A1', 'CPV').json()['id']
    sessions = client.container.sessions
    results, barrier = [], threading.Barrier(12)

    def worker():
        barrier.wait()
        results.append(sessions.record(sid, 'S1', 'An', 0.9)['status'])

    threads = [threading.Thread(target=worker) for _ in range(12)]
    [t.start() for t in threads]
    [t.join() for t in threads]
    assert results.count('inserted') == 1 and results.count('already') == 11
    with sqlite3.connect(client.container.paths.db_path) as db:
        assert db.execute('SELECT COUNT(*) FROM attendance').fetchone()[0] == 1


def test_late_result_after_end_is_rejected(h, client):
    h.course('A1', 'CPV')
    _, e = h.enroll('A1', 'S1', 'An')
    sid = h.open('A1', 'CPV').json()['id']
    client.post(f'/api/sessions/{sid}/end')
    res = h.frame(sid, h.face_frame(e))
    assert res['status'] == 'session_closed'
    # a recognition that was already in flight when the session ended
    assert client.container.sessions.record(sid, 'S1', 'An', 0.9)['status'] == 'session_closed'
    with sqlite3.connect(client.container.paths.db_path) as db:
        assert db.execute('SELECT COUNT(*) FROM attendance').fetchone()[0] == 0


def test_deadline_expiry_rejects_results(h, client, monkeypatch):
    h.course('A1', 'CPV')
    _, e = h.enroll('A1', 'S1', 'An')
    sid = h.open('A1', 'CPV', minutes=1).json()['id']
    later = clock.now() + timedelta(minutes=2)
    monkeypatch.setattr(clock, 'now', lambda: later)
    assert client.get(f'/api/sessions/{sid}').json()['state'] == 'expired'
    assert h.frame(sid, h.face_frame(e))['status'] == 'session_closed'
    client.container.tick()
    assert client.get(f'/api/sessions/{sid}').json()['end_reason'] == 'timeout'


def test_session_across_midnight_keeps_its_date(h, client, monkeypatch):
    h.course('A1', 'CPV')
    _, e = h.enroll('A1', 'S1', 'An')
    t0 = datetime(2026, 10, 22, 23, 58, 0)
    monkeypatch.setattr(clock, 'now', lambda: t0)
    sid = h.open('A1', 'CPV', minutes=None).json()['id']
    monkeypatch.setattr(clock, 'now', lambda: t0 + timedelta(minutes=5))
    assert h.frame(sid, h.face_frame(e))['status'] == 'recognized'
    with sqlite3.connect(client.container.paths.db_path) as db:
        date, ts = db.execute('SELECT date, timestamp FROM attendance').fetchone()
    assert date == '2026-10-22' and ts.startswith('2026-10-23 00:03')


# ---------------------------------------------------------------- students
def test_lowercase_id_and_leading_zero(h, client):
    h.course('A1', 'CPV')
    r, _ = h.enroll('A1', '  he0012 ', 'An')
    assert r.status_code == 201 and r.json()['student_id'] == 'HE0012'
    r, _ = h.enroll('A1', '000123', 'Bình')
    assert r.status_code == 201 and r.json()['student_id'] == '000123'
    r, _ = h.enroll('A1', 'HE0012', 'Chi')
    assert r.status_code == 409 and r.json()['code'] == 'duplicate_id'
    folders = sorted(os.listdir(os.path.join(client.container.paths.dataset_dir, 'A1')))
    assert 'A1_HE0012_An' in folders and 'A1_000123_Bình' in folders


def test_vietnamese_name_and_spaces(h, client):
    h.course('A1', 'CPV')
    nfd = unicodedata.normalize('NFD', '  Nguyễn   Văn  Ánh ')
    r, e = h.enroll('A1', 'S9', nfd)
    st = r.json()
    assert r.status_code == 201 and st['full_name'] == 'Nguyễn Văn Ánh'
    assert st['folder'] == 'A1_S9_Nguyễn_Văn_Ánh'
    sid = h.open('A1', 'CPV').json()['id']
    res = h.frame(sid, h.face_frame(e))
    assert res['student']['student_name'] == 'Nguyễn Văn Ánh'
    assert client.post('/api/students', data={'class_name': 'A1', 'student_id': 'X', 'full_name': 'a_b'},
                       files=[('files', ('a.jpg', h.face_frame(h.new_embedding()), 'image/jpeg'))]).status_code == 422


def test_edit_student_does_not_touch_other_class(h, client):
    h.course('A1', 'CPV')
    h.course('B1', 'CPV')
    ra, ea = h.enroll('A1', 'S1', 'An')
    _, eb = h.enroll('B1', 'S1', 'Bình')
    s1 = h.open('A1', 'CPV').json()['id']
    s2 = h.open('B1', 'CPV').json()['id']
    h.frame(s1, h.face_frame(ea))
    h.frame(s2, h.face_frame(eb))
    r = client.patch(f"/api/students/{ra.json()['id']}", json={'student_id': 's1x', 'full_name': 'An  Nguyễn'})
    assert r.status_code == 200 and r.json()['student_id'] == 'S1X'
    with sqlite3.connect(client.container.paths.db_path) as db:
        rows = dict(db.execute('SELECT class_name, student_id || "|" || student_name FROM attendance'))
    assert rows == {'A1': 'S1X|An Nguyễn', 'B1': 'S1|Bình'}
    # the renamed student is still recognised under the new identity
    assert h.frame(s1, h.face_frame(ea))['student']['student_id'] == 'S1X'


def test_desktop_update_student_info_is_class_scoped(tmp_path):
    from functions.database import Database
    db = Database(str(tmp_path / 'a.db'))
    db.save_attendance('A1', 'CPV', 'S1', 'An', 0.9, slot=1)
    db.save_attendance('B1', 'CPV', 'S1', 'Bình', 0.9, slot=1)
    db.update_student_info('S1', 'A1', 'S2', 'An mới')
    with sqlite3.connect(db.db_path) as conn:
        rows = sorted(conn.execute('SELECT class_name, student_id, student_name FROM attendance'))
    assert rows == [('A1', 'S2', 'An mới'), ('B1', 'S1', 'Bình')]


def test_duplicate_face_rejected_on_enroll(h):
    h.course('A1', 'CPV')
    _, e = h.enroll('A1', 'S1', 'An')
    r, _ = h.enroll('A1', 'S2', 'Giả', emb=e)
    assert r.status_code == 409 and r.json()['code'] == 'duplicate_face'


def test_enroll_sample_errors(h, client):
    h.course('A1', 'CPV')
    no_face = h.fake.scene([])
    two = h.fake.scene([((50, 50, 150, 180), h.new_embedding()), ((350, 60, 150, 180), h.new_embedding())])
    tiny = h.fake.scene([((10, 10, 30, 30), h.new_embedding())])
    bad_emb = h.fake.scene([((200, 100, 200, 240), np.full(512, np.nan))])
    for jpg, reason in ((no_face, 'no_face'), (two, 'multiple_faces'), (tiny, 'too_small'), (bad_emb, 'bad_embedding')):
        r = client.post('/api/enrollment/check', files={'file': ('a.jpg', jpg, 'image/jpeg')})
        assert r.status_code == 200 and r.json()['reason'] == reason
    r = client.post('/api/students', data={'class_name': 'A1', 'student_id': 'S1', 'full_name': 'An'},
                    files=[('files', ('a.jpg', two, 'image/jpeg'))])
    assert r.status_code == 422 and r.json()['code'] == 'invalid_samples'
    assert client.get('/api/students').json() == []


def test_delete_student_keeps_history_and_files(h, client):
    h.course('A1', 'CPV')
    r, e = h.enroll('A1', 'S1', 'An')
    sid = h.open('A1', 'CPV').json()['id']
    h.frame(sid, h.face_frame(e))
    assert client.delete(f"/api/students/{r.json()['id']}").status_code == 200
    assert h.frame(sid, h.face_frame(e))['status'] in ('no_students', 'unknown')
    trash = os.path.join(client.container.paths.dataset_dir, 'A1', '.deleted')
    assert len(os.listdir(trash)) == 1
    assert client.get('/api/history').json()['summary']['present'] == 1


# ---------------------------------------------------------------- recognition rules
def test_enroll_and_recognition_use_adaface_not_buffalo_arcface(h, client):
    """buffalo_sc attaches its own ArcFace vector to every face; it must never reach the AdaFace gallery."""
    import json
    h.course('A1', 'CPV')
    r, e = h.enroll('A1', 'S1', 'An')
    folder = os.path.join(client.container.paths.dataset_dir, 'A1', r.json()['folder'])
    stored = np.load(os.path.join(folder, 'embeddings.npy'))
    assert float((stored @ e).min()) > 0.8            # AdaFace vectors, not the random ArcFace decoys
    meta = json.load(open(os.path.join(folder, 'model_meta.json'), encoding='utf-8'))
    assert meta['model_id'] == 'cvlface_adaface_ir101_webface12m'
    ada = client.container.models.adaface
    before = ada.calls
    sid = h.open('A1', 'CPV').json()['id']
    assert h.frame(sid, h.face_frame(e))['status'] == 'recognized'
    assert ada.calls > before


def test_tracking_reuses_adaface_result_while_face_stays(h, client):
    h.course('A1', 'CPV')
    _, e = h.enroll('A1', 'S1', 'An')
    sid = h.open('A1', 'CPV').json()['id']
    rec = client.container.recognition
    now = [1000.0]
    rec._clock = lambda: now[0]
    ada = client.container.models.adaface
    assert h.frame(sid, h.face_frame(e))['status'] == 'recognized'
    calls = ada.calls
    r = h.frame(sid, h.face_frame(e, box=(212, 108, 200, 240)))      # same face, moved a little
    assert r['status'] == 'already_recorded' and r['cached'] and ada.calls == calls
    assert r['faces'][0]['state'] == 'recognized' and r['student']['student_id'] == 'S1'
    now[0] += 3                                                        # verify interval elapsed
    r = h.frame(sid, h.face_frame(e))
    assert not r.get('cached') and ada.calls == calls + 1


def test_tracking_resets_when_face_leaves(h, client):
    h.course('A1', 'CPV')
    _, e = h.enroll('A1', 'S1', 'An')
    sid = h.open('A1', 'CPV').json()['id']
    client.container.recognition._clock = lambda: 1000.0               # cache would never expire
    assert h.frame(sid, h.face_frame(e))['status'] == 'recognized'
    assert h.frame(sid, h.fake.scene([]))['status'] == 'no_face'
    stranger = h.frame(sid, h.face_frame(h.new_embedding()))           # same spot, new person
    assert stranger['status'] == 'unknown' and not stranger.get('cached')


def test_face_without_landmarks_is_unverifiable():
    from functions.adaface_service import embed_detected_face

    class NoKps:
        kps = None
        embedding = np.ones(512, np.float32)       # an ArcFace vector alone is not enough

    with pytest.raises(ValueError):
        embed_detected_face(object(), np.zeros((10, 10, 3), np.uint8), NoKps())


def test_stranger_rejected(h, client):
    h.course('A1', 'CPV')
    h.enroll('A1', 'S1', 'An')
    sid = h.open('A1', 'CPV').json()['id']
    res = h.frame(sid, h.face_frame(h.new_embedding()))
    assert res['status'] == 'unknown' and res['faces'][0]['state'] == 'unknown'
    assert client.get(f'/api/sessions/{sid}').json()['present_count'] == 0


def test_ambiguous_identity_not_recorded(h, client):
    h.course('A1', 'CPV')
    h.enroll('A1', 'S1', 'An')
    rec = client.container.models.recognizer('A1')
    stored = unit(rec.embeddings_db['A1_S1_An'][0])
    other = unit(stored + 0.35 * h.new_embedding())
    rec.embeddings_db['A1_S2_Bình'] = np.atleast_2d(other)   # look-alike enrolled directly
    probe = unit(stored + other)                              # equally close to both
    sid = h.open('A1', 'CPV').json()['id']
    res = h.frame(sid, h.face_frame(probe))
    assert res['status'] == 'ambiguous'
    assert client.get(f'/api/sessions/{sid}').json()['present_count'] == 0


def test_multiple_faces_rules(h, client):
    h.course('A1', 'CPV')
    _, e = h.enroll('A1', 'S1', 'An')
    sid = h.open('A1', 'CPV').json()['id']
    blocker = ((420, 120, 160, 200), h.new_embedding())
    res = h.frame(sid, h.face_frame(e, extra=[blocker]))
    assert res['status'] == 'multiple_faces' and {f['state'] for f in res['faces']} == {'multiple'}
    assert client.get(f'/api/sessions/{sid}').json()['present_count'] == 0
    bystander = ((600, 10, 20, 20), h.new_embedding())    # tiny, far away → ignored
    res = h.frame(sid, h.face_frame(e, extra=[bystander]))
    assert res['status'] == 'recognized'
    assert [f['state'] for f in res['faces']] == ['recognized', 'ignored']


def test_preview_mode_does_not_record(h, client):
    h.course('A1', 'CPV')
    _, e = h.enroll('A1', 'S1', 'An')
    sid = h.open('A1', 'CPV').json()['id']
    assert h.frame(sid, h.face_frame(e), record=False)['status'] == 'preview'
    assert client.get(f'/api/sessions/{sid}').json()['present_count'] == 0


# ---------------------------------------------------------------- anti-spoof
def test_antispoof_errors_are_never_live():
    from functions.anti_spoof import AntiSpoof
    spoof = AntiSpoof('/nonexistent')
    frame = np.zeros((100, 100, 3), np.uint8)
    assert spoof.predict(None, (0, 0, 10, 10))['is_real'] is False
    assert spoof.predict(frame, (500, 500, 10, 10))['is_real'] is False          # empty ROI
    assert spoof.predict(frame, ('x', 0, 1, 1))['is_real'] is False             # exception path
    res = spoof.predict(frame, (10, 10, 50, 50), allow_texture=False)           # no model
    assert res['is_real'] is False and res['label'] == 'ERROR'


def test_antispoof_model_mode_without_model_blocks(h, client):
    h.course('A1', 'CPV')
    _, e = h.enroll('A1', 'S1', 'An')
    client.put('/api/settings', json={'antispoof_mode': 'model'})
    assert client.get('/api/models/status').json()['antispoof']['state'] == 'missing_model'
    sid = h.open('A1', 'CPV').json()['id']
    res = h.frame(sid, h.face_frame(e))
    assert res['status'] == 'spoof_error' and res['faces'][0]['state'] == 'unverifiable'
    assert client.get(f'/api/sessions/{sid}').json()['present_count'] == 0


def test_antispoof_inference_error_blocks(h, client, monkeypatch):
    h.course('A1', 'CPV')
    _, e = h.enroll('A1', 'S1', 'An')
    client.put('/api/settings', json={'antispoof_mode': 'texture_experimental'})
    spoof = client.container.models.anti_spoof
    monkeypatch.setattr(spoof, '_predict_cv', lambda roi: (_ for _ in ()).throw(RuntimeError('boom')))
    sid = h.open('A1', 'CPV').json()['id']
    assert h.frame(sid, h.face_frame(e))['status'] == 'spoof_error'
    monkeypatch.setattr(spoof, '_predict_cv', lambda roi: {'is_real': False, 'live_score': 0.1, 'label': 'FAKE'})
    assert h.frame(sid, h.face_frame(e))['status'] == 'spoof_rejected'
    monkeypatch.setattr(spoof, '_predict_cv', lambda roi: {'is_real': True, 'live_score': 0.9, 'label': 'LIVE'})
    res = h.frame(sid, h.face_frame(e))
    assert res['status'] == 'recognized' and res['antispoof']['experimental']


# ---------------------------------------------------------------- invalid input / missing parts
def test_invalid_and_oversized_frames(h, client):
    h.course('A1', 'CPV')
    sid = h.open('A1', 'CPV').json()['id']
    assert h.frame(sid, b'not an image')['status'] == 'invalid_frame'
    r = client.post(f'/api/sessions/{sid}/frame', files={'file': ('big.jpg', b'0' * (3 * 1024 * 1024 + 10), 'image/jpeg')})
    assert r.status_code == 413
    assert client.post('/api/sessions/9999/frame', files={'file': ('f.jpg', b'x', 'image/jpeg')}).status_code == 404


def test_missing_model_reports_unavailable(data_dir, monkeypatch):
    import functions.face_recognizer as fr
    monkeypatch.setattr(fr, '_insight_app', None)
    monkeypatch.setattr(fr, 'FaceAnalysis', lambda **kw: (_ for _ in ()).throw(FileNotFoundError('buffalo_sc missing')))
    c = make_container(data_dir, load=False)
    c.models._load()
    assert c.models.state == 'error' and 'buffalo_sc' in c.models.error
    c.catalog.create_class('A1')
    c.catalog.create_course('A1', 'CPV')
    sid = c.sessions.open('A1', 'CPV', 1, 10)['id']
    import cv2
    jpg = cv2.imencode('.jpg', np.zeros((120, 160, 3), np.uint8))[1].tobytes()
    assert c.recognition.process_frame(sid, jpg)['status'] == 'model_unavailable'


def test_websocket_frame_roundtrip(h, client):
    h.course('A1', 'CPV')
    _, e = h.enroll('A1', 'S1', 'An')
    sid = h.open('A1', 'CPV').json()['id']
    jpg = h.face_frame(e)
    with client.websocket_connect(f'/ws/sessions/{sid}') as ws:
        ws.send_bytes((7).to_bytes(4, 'little') + bytes([1]) + jpg)
        res = ws.receive_json()
    assert res['frame_id'] == 7 and res['session_id'] == sid and res['status'] == 'recognized'


def test_rtsp_validation_and_masking(client, h):
    from backend.services.settings import mask_rtsp
    assert mask_rtsp('rtsp://admin:s3cret@10.0.0.2:554/s') == 'rtsp://admin:***@10.0.0.2:554/s'
    assert 's3cret' not in mask_rtsp('rtsp://s3cret@10.0.0.2/s')
    r = client.put('/api/settings', json={'rtsp_url': 'http://x'})
    assert r.status_code == 422
    r = client.put('/api/settings', json={'rtsp_url': 'rtsp://admin:s3cret@10.0.0.2:554/s'})
    assert r.status_code == 200 and 's3cret' not in r.text and r.json()['has_rtsp_url']
    h.course('A1', 'CPV')
    sid = h.open('A1', 'CPV').json()['id']
    assert client.post(f'/api/sessions/{sid}/rtsp/start', json={'url': 'not-rtsp'}).status_code == 422


def test_settings_validation_and_persistence(client, data_dir, fake_app):
    assert client.put('/api/settings', json={'cosine_threshold': 1.5}).status_code == 422
    assert client.put('/api/settings', json={'antispoof_mode': 'magic'}).status_code == 422
    r = client.put('/api/settings', json={'cosine_threshold': 0.7, 'min_margin': 0.1, 'preprocessing': True})
    assert r.status_code == 200
    again = make_container(data_dir, load=False)
    s = again.settings.get()
    assert (s['cosine_threshold'], s['min_margin'], s['preprocessing']) == (0.7, 0.1, True)


# ---------------------------------------------------------------- reports / persistence
def test_excel_export_matches_sqlite(h, client):
    h.course('A1', 'CPV')
    _, e1 = h.enroll('A1', '007', 'Nguyễn Văn An')
    h.enroll('A1', 'S2', 'Trần Bình')
    sid = h.open('A1', 'CPV', slot=4).json()['id']
    h.frame(sid, h.face_frame(e1))
    client.post(f'/api/sessions/{sid}/end')
    hist = client.get('/api/history?class_name=A1').json()
    r = client.get('/api/reports/export?class_name=A1')
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    rows = list(wb['Chi tiết'].iter_rows(min_row=2, values_only=True))
    assert len(rows) == len(hist['records']) == 2
    by_id = {row[5]: row for row in rows}
    assert by_id['007'][7] == 'Có mặt' and by_id['S2'][7] == 'Vắng'   # leading zero preserved as text
    with sqlite3.connect(client.container.paths.db_path) as db:
        ts, conf = db.execute("SELECT timestamp, confidence FROM attendance WHERE student_id='007'").fetchone()
    assert by_id['007'][8] == ts and by_id['007'][9] == pytest.approx(conf, abs=1e-4)
    # desktop-style class sheet was synced from SQLite when the session ended
    import time
    path = os.path.join(client.container.paths.dataset_dir, 'A1', 'attendance.xlsx')
    for _ in range(50):
        if os.path.exists(path) and len(next(load_workbook(path).active.iter_rows(max_row=1, values_only=True))) > 2:
            break
        time.sleep(0.1)
    sheet = load_workbook(path).active
    header = next(sheet.iter_rows(max_row=1, values_only=True))
    assert any(str(c).startswith('A1 - CPV - Slot 4') for c in header)


def test_absent_counts_use_session_roster(h, client):
    h.course('A1', 'CPV')
    h.course('B1', 'CPV')
    _, e = h.enroll('A1', 'S1', 'An')
    h.enroll('A1', 'S2', 'Bình')
    h.enroll('B1', 'T1', 'Chi')          # other class must not count as absent in A1
    sid = h.open('A1', 'CPV').json()['id']
    h.frame(sid, h.face_frame(e))
    h.enroll('A1', 'S3', 'Dung')          # enrolled after the session ended? no: still open → joins roster
    client.post(f'/api/sessions/{sid}/end')
    h.enroll('A1', 'S4', 'Em')            # after the end → must not change the old report
    day = client.get('/api/dashboard').json()['day']
    assert (day['present'], day['absent'], day['roster']) == (1, 2, 3)


def test_restart_keeps_data(h, client, data_dir, fake_app):
    h.course('A1', 'CPV')
    _, e = h.enroll('A1', 'S1', 'An')
    sid = h.open('A1', 'CPV').json()['id']
    h.frame(sid, h.face_frame(e))
    c2 = make_container(data_dir)
    assert [x['name'] for x in c2.catalog.list_classes()] == ['A1']
    assert len(c2.students.list()) == 1
    s = c2.sessions.get(sid)
    assert s['present_count'] == 1 and s['state'] == 'active'
    assert c2.recognition.process_frame(sid, h.face_frame(e))['status'] == 'already_recorded'


def test_migration_of_legacy_database(tmp_path, fake_app, monkeypatch):
    data = tmp_path / 'legacy'
    (data / 'dataset' / 'SE1').mkdir(parents=True)
    db = sqlite3.connect(data / 'attendance.db')
    db.execute('CREATE TABLE attendance (id INTEGER PRIMARY KEY AUTOINCREMENT, class_name TEXT NOT NULL, '
               'subject TEXT NOT NULL, student_id TEXT NOT NULL, student_name TEXT NOT NULL, timestamp TEXT NOT NULL, '
               'date TEXT NOT NULL, confidence REAL DEFAULT 0.0)')
    rows = [('SE1', 'CPV', 'S1', 'An', '2026-06-29 07:52:15', '2026-06-29', 0.9),
            ('SE1', 'CPV', 'S2', 'Bình', '2026-06-29 07:53:00', '2026-06-29', 0.8),
            ('SE1', 'CPV', 'S2', 'Bình', '2026-06-29 07:54:00', '2026-06-29', 0.8)]   # legacy duplicate
    db.executemany('INSERT INTO attendance(class_name, subject, student_id, student_name, timestamp, date, confidence) '
                   'VALUES (?, ?, ?, ?, ?, ?, ?)', rows)
    db.commit()
    db.close()
    (data / 'sessions.json').write_text('[{"class": "SE1", "subject": "CPV"}, {"class": "bad_name", "subject": "X"}]',
                                        encoding='utf-8')
    monkeypatch.setenv('SA_FRONTEND_DIST', str(tmp_path / 'nodist'))
    c = make_container(str(data))
    rep = c.migration
    assert rep['backup'] and os.path.exists(rep['backup'])
    assert rep['legacy_sessions'] == 1 and rep['imported_courses'] == 1 and rep['skipped_courses']
    assert rep['duplicates'] and not rep['unique_index']           # reported, nothing deleted
    with sqlite3.connect(data / 'attendance.db') as conn:
        assert conn.execute('SELECT COUNT(*) FROM attendance').fetchone()[0] == 3
    hist = c.reports.history({})
    assert hist['summary']['legacy_sessions'] == 1 and hist['summary']['absent'] == 0
    assert {r['status'] for r in hist['records']} == {'present'}
    # second run is a no-op
    c2 = make_container(str(data))
    assert c2.migration['backup'] is None and c2.migration['legacy_sessions'] == 0


def test_frontend_is_served_when_built(tmp_path, data_dir, fake_app, monkeypatch):
    from fastapi.testclient import TestClient
    from backend.app import create_app
    dist = tmp_path / 'dist'
    (dist / 'assets').mkdir(parents=True)
    (dist / 'index.html').write_text('<!doctype html><div id="root"></div>', encoding='utf-8')
    (dist / 'assets' / 'app.js').write_text('console.log(1)', encoding='utf-8')
    monkeypatch.setenv('SA_FRONTEND_DIST', str(dist))
    c = make_container(data_dir)
    with TestClient(create_app(c, load_models=False)) as cl:
        assert 'id="root"' in cl.get('/').text
        assert 'id="root"' in cl.get('/students/enroll').text          # SPA route
        assert cl.get('/assets/app.js').status_code == 200
        assert cl.get('/api/health').json()['ok'] is True
        assert cl.get('/api/nope').status_code == 404
